from pathlib import Path

import pytest

from bank_statement_processor import BankStatementProcessor, Transaction

SAMPLE_TEXT = """
HDFC BANK LIMITED
Account Holder Name: Jane Doe
Account No: 123456789012
IFSC: HDFC0001234

Date        Description                          Debit      Credit      Balance
03/09/2024  SALARY CREDIT                       0.00       50000.00   250000.00
05/09/2024  AMAZON PAY                           2999.00    0.00       247001.00
07/09/2024  PETROL PUMP                          1500.00    0.00       245501.00
12/09/2024  RENT PAYMENT                         22000.00   0.00       223501.00
12/09/2024  RENT PAYMENT                         22000.00   0.00       223501.00
15/09/2024  Opening Balance                      0.00       0.00       223501.00
"""


def test_pdf_text_extraction_parses_transactions_and_metadata() -> None:
    processor = BankStatementProcessor()
    statement = processor.process_text(SAMPLE_TEXT, "sample.pdf")

    assert statement.account_holder_name == "Jane Doe"
    assert statement.account_number == "123456789012"
    assert statement.ifsc_code == "HDFC0001234"
    assert statement.bank_name and "HDFC" in statement.bank_name.upper()
    assert statement.document_type == "text"
    assert len(statement.transactions) == 4
    assert statement.transactions[0].description == "SALARY CREDIT"
    assert statement.transactions[0].credit_amount == 50000.0
    assert statement.transactions[1].category == "Shopping"
    assert statement.transactions[2].category == "Transport"
    assert statement.transactions[3].category == "Housing"


def test_cr_dr_layout_and_wrapped_narration() -> None:
    text = """
    ICICI Bank
    Customer Name: Rohan Mehta
    A/C No: 998877665544
    IFSC CODE: ICIC0000777

    Date Particulars Amount Balance
    19-Sep-2024 UPI/SWIGGY/FOOD 450.00 DR 12000.00
    late night order
    20-Sep-2024 NEFT SALARY ACME CORP 75000.00 CR 87000.00
    """
    statement = BankStatementProcessor().process_text(text, "icici.txt")
    assert statement.account_holder_name == "Rohan Mehta"
    assert len(statement.transactions) == 2
    assert "late night order" in statement.transactions[0].description
    assert statement.transactions[0].debit_amount == 450.0
    assert statement.transactions[0].category == "Food"
    assert statement.transactions[1].credit_amount == 75000.0
    assert statement.transactions[1].category == "Salary"


def test_rule_based_classification_handles_common_categories() -> None:
    transactions = [
        Transaction(date="2024-09-03", description="SALARY CREDIT", debit_amount=0.0, credit_amount=50000.0, balance=120000.0),
        Transaction(date="2024-09-05", description="NETFLIX SUBSCRIPTION", debit_amount=599.0, credit_amount=0.0, balance=119401.0),
        Transaction(date="2024-09-06", description="ELECTRICITY BILL", debit_amount=1250.0, credit_amount=0.0, balance=118151.0),
        Transaction(date="2024-09-08", description="ATM CASH WITHDRAWAL", debit_amount=2000.0, credit_amount=0.0, balance=116151.0),
    ]

    classified = BankStatementProcessor().classify_transactions(transactions)

    assert [item.category for item in classified] == [
        "Salary",
        "Entertainment",
        "Bills",
        "Cash Withdrawal",
    ]
    assert all(item.classification_method == "heuristic" for item in classified)


def test_machine_learning_classifies_unseen_merchant_wording() -> None:
    classified = BankStatementProcessor().classify_transactions(
        [Transaction(date="2024-09-01", description="GROWW SIP MUTUAL FUND PURCHASE", debit_amount=5000.0)]
    )
    assert classified[0].category == "Investment"
    assert classified[0].classification_method in {"heuristic", "machine_learning"}
    assert classified[0].confidence > 0


def test_exporter_writes_csv_and_xlsx(tmp_path: Path) -> None:
    processor = BankStatementProcessor()
    statement = processor.process_text(SAMPLE_TEXT, "sample.pdf")

    csv_path = tmp_path / "bank_statement.csv"
    xlsx_path = tmp_path / "bank_statement.xlsx"

    processor.export_transactions(statement, csv_path)
    processor.export_transactions(statement, xlsx_path)

    csv_text = csv_path.read_text(encoding="utf-8")
    assert csv_path.exists()
    assert xlsx_path.exists()
    assert "SALARY CREDIT" in csv_text
    assert "category" in csv_text
    assert xlsx_path.stat().st_size > 0


def test_empty_statement_raises() -> None:
    processor = BankStatementProcessor()
    try:
        processor.process_text("   \n  ", "blank.txt")
    except ValueError as error:
        assert "empty" in str(error).lower()
    else:
        raise AssertionError("Expected empty statements to raise ValueError")


def test_extracts_split_lines_and_amounts_without_decimals() -> None:
    text = """
    SBI BANK
    Account Holder Name: Asha Rao
    Account Number: 445566778899
    IFSC: SBIN0001122

    Txn Date  Narration  Withdrawal  Deposit  Balance
    01/09/24
    UPI-AMAZON PAY INDIA
    1999
    0
    12001
    02-09-2024  NEFT SALARY ACME  0  45000  57001
    """
    statement = BankStatementProcessor().process_text(text, "sbi.pdf")
    assert statement.account_holder_name == "Asha Rao"
    assert len(statement.transactions) >= 2
    amazon = statement.transactions[0]
    assert "AMAZON" in amazon.description.upper()
    assert amazon.debit_amount == 1999.0
    assert statement.transactions[1].credit_amount == 45000.0
    assert statement.transactions[1].category == "Salary"


def test_reconstructed_column_layout_keeps_row_alignment() -> None:
    fitz = pytest.importorskip("fitz")
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 60), "HDFC BANK LIMITED")
    page.insert_text((72, 78), "Account Name: Jane Doe")
    page.insert_text((72, 96), "Account No: 123456789012")
    page.insert_text((72, 114), "IFSC: HDFC0001234")
    rows = [
        ("03/09/2024", "SALARY CREDIT", "0.00", "50000.00", "250000.00"),
        ("05/09/2024", "AMAZON PAY", "2999.00", "0.00", "247001.00"),
    ]
    for index, (date, desc, debit, credit, balance) in enumerate(rows):
        y = 150 + index * 18
        page.insert_text((72, y), date)
        page.insert_text((160, y), desc)
        page.insert_text((320, y), debit)
        page.insert_text((400, y), credit)
        page.insert_text((480, y), balance)
    pdf_bytes = document.tobytes()
    document.close()

    statement = BankStatementProcessor().process_pdf(pdf_bytes, "columns.pdf")
    assert len(statement.transactions) >= 2
    assert statement.transactions[0].credit_amount == 50000.0
    assert statement.transactions[1].category == "Shopping"


def test_process_pdf_accepts_generated_statement() -> None:
    fitz = pytest.importorskip("fitz")
    text = SAMPLE_TEXT.strip()
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    pdf_bytes = document.tobytes()
    document.close()

    statement = BankStatementProcessor().process_pdf(pdf_bytes, "generated.pdf")
    assert statement.document_type == "text"
    assert statement.account_number == "123456789012"
    assert len(statement.transactions) >= 4
    assert statement.transactions[0].category == "Salary"


CSV_ROWS = [
    ["Date", "Description", "Debit", "Credit", "Balance"],
    ["03/09/2024", "SALARY CREDIT", "0.00", "50000.00", "250000.00"],
    ["05/09/2024", "AMAZON PAY", "2999.00", "0.00", "247001.00"],
    ["07/09/2024", "PETROL PUMP", "1500.00", "0.00", "245501.00"],
]


def test_process_tabular_csv_parses_transactions_and_classifies() -> None:
    processor = BankStatementProcessor()
    statement = processor.process_tabular(CSV_ROWS, "statement.csv")

    assert statement.document_type == "tabular"
    assert len(statement.transactions) == 3
    assert statement.transactions[0].credit_amount == 50000.0
    assert statement.transactions[0].category == "Salary"
    assert statement.transactions[0].confidence > 0
    assert statement.transactions[1].category == "Shopping"
    assert statement.transactions[2].category == "Transport"


def test_process_tabular_debit_credit_without_balance() -> None:
    rows = [
        ["Value Date", "Narration", "Debit", "Credit"],
        ["03/09/2024", "UPI/SWIGGY/FOOD", "450.00", "0.00"],
        ["04/09/2024", "NEFT SALARY ACME", "0.00", "75000.00"],
    ]
    statement = BankStatementProcessor().process_tabular(rows, "no-balance.csv")

    assert len(statement.transactions) == 2
    assert statement.transactions[0].debit_amount == 450.0
    assert statement.transactions[0].category == "Food"
    assert statement.transactions[1].credit_amount == 75000.0
    assert statement.transactions[1].category == "Salary"


def test_process_tabular_excel_export(tmp_path: Path) -> None:
    from openpyxl import Workbook, load_workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Txn Date", "Narration", "Withdrawal", "Deposit", "Balance"])
    sheet.append(["03/09/2024", "ATM CASH WITHDRAWAL", "2000.00", "0.00", "248000.00"])
    sheet.append(["05/09/2024", "INTEREST CREDIT", "0.00", "1250.00", "249250.00"])
    path = tmp_path / "statement.xlsx"
    workbook.save(path)

    loaded = load_workbook(path, read_only=True, data_only=True)
    rows = [list(row) for sheet in loaded.worksheets for row in sheet.iter_rows(values_only=True)]
    loaded.close()

    statement = BankStatementProcessor().process_tabular(rows, "statement.xlsx")

    assert len(statement.transactions) == 2
    assert statement.transactions[0].debit_amount == 2000.0
    assert statement.transactions[0].category == "Cash Withdrawal"
    assert statement.transactions[1].credit_amount == 1250.0
    assert statement.transactions[1].category == "Interest"


def test_process_image_statement_uses_ocr(monkeypatch) -> None:
    from io import BytesIO

    from PIL import Image

    buffer = BytesIO()
    Image.new("RGB", (120, 40), "white").save(buffer, format="PNG")

    processor = BankStatementProcessor()
    monkeypatch.setattr(processor, "_ocr_image", lambda image: SAMPLE_TEXT)

    statement = processor.process_image(buffer.getvalue(), "statement.png")

    assert statement.document_type == "image"
    assert statement.account_number == "123456789012"
    assert len(statement.transactions) == 4


def test_process_image_statement_raises_when_invalid() -> None:
    processor = BankStatementProcessor()

    try:
        processor.process_image(b"", "blank.png")
    except ValueError as error:
        assert "empty" in str(error).lower()
    else:
        raise AssertionError("Expected empty image bytes to raise ValueError")
