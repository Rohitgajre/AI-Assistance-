"""CSV and XLSX extraction with row-level context."""

import csv
from io import BytesIO, StringIO

from document_analysis.exceptions import InvalidDocumentError
from document_analysis.schemas import DocumentChunk, SourceReference


def extract_spreadsheet(content: bytes, extension: str) -> list[DocumentChunk]:
    if extension == "csv":
        try:
            rows = list(csv.reader(StringIO(content.decode("utf-8-sig", errors="replace"))))
        except csv.Error as error:
            raise InvalidDocumentError("The CSV file could not be read.") from error
        return _rows_to_chunks(rows, "CSV")
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        chunks = []
        for sheet in workbook.worksheets:
            chunks.extend(_rows_to_chunks(list(sheet.iter_rows(values_only=True)), sheet.title))
        if not chunks:
            raise InvalidDocumentError("The spreadsheet is empty.")
        return chunks
    except ImportError as error:
        raise InvalidDocumentError("XLSX support is unavailable; install openpyxl.") from error
    except InvalidDocumentError:
        raise
    except Exception as error:
        raise InvalidDocumentError("The XLSX file could not be read.") from error


def _rows_to_chunks(rows: list[object], section: str) -> list[DocumentChunk]:
    chunks = []
    for index, row in enumerate(rows, start=1):
        values = row if isinstance(row, (list, tuple)) else [row]
        text = " | ".join(str(value) for value in values if value is not None).strip()
        if text:
            chunks.append(
                DocumentChunk(
                    text=text,
                    source=SourceReference(section=f"{section}, row {index}", excerpt=text[:300]),
                )
            )
    if not chunks:
        raise InvalidDocumentError("The spreadsheet is empty.")
    return chunks
