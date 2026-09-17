"""PPTX slide text extraction."""

from io import BytesIO

from document_analysis.exceptions import InvalidDocumentError
from document_analysis.schemas import DocumentChunk, SourceReference


def extract_pptx(content: bytes) -> list[DocumentChunk]:
    try:
        from pptx import Presentation

        chunks = []
        for number, slide in enumerate(Presentation(BytesIO(content)).slides, start=1):
            text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")).strip()
            if text:
                chunks.append(
                    DocumentChunk(
                        text=text,
                        source=SourceReference(section=f"slide {number}", excerpt=text[:300]),
                    )
                )
        if not chunks:
            raise InvalidDocumentError("No readable text was found in the presentation.")
        return chunks
    except ImportError as error:
        raise InvalidDocumentError("PPTX support is unavailable; install python-pptx.") from error
    except InvalidDocumentError:
        raise
    except Exception as error:
        raise InvalidDocumentError("The PPTX file could not be read.") from error
