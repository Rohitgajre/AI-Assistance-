"""Streamlit entry point for Nebula AI — ChatGPT-style chat with intelligent document handling."""
from __future__ import annotations

import hashlib
import logging
import os
import re
import threading
import uuid
from collections.abc import Iterator
from datetime import datetime
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Any, Final

import streamlit as st
from dotenv import load_dotenv
from langchain_google_genai import ChatGoogleGenerativeAI

from agent.orchestrator import get_orchestrator
from bank_statement_processor import BankStatement, BankStatementProcessor, Transaction
from chat_service import ask_model
from document_analysis.config import OllamaSettings
from document_analysis.exceptions import DocumentAnalysisError
from document_analysis.service import DocumentAnalysisService, OllamaClient
from ingestion.document_router import route_document

# New infrastructure imports
from ingestion.service import get_ingestion_service
from ingestion.validators import get_extension, validate_upload
from memory.conversation_store import Conversation, Message, get_conversation_store
from memory.document_registry import DocumentRecord, get_document_registry
from providers.embedding_provider import get_embedding_provider
from rag.chunking import StructureAwareChunker
from rag.vector_store import get_vector_store

load_dotenv()

# Setup comprehensive logging
LOG_DIR = Path(os.getenv("NEBULA_LOG_DIR") or "./logs")
LOG_DIR.mkdir(parents=True, exist_ok=True)
log_file = LOG_DIR / "nebula_ai.log"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s - %(message)s",
    handlers=[
        RotatingFileHandler(log_file, maxBytes=10*1024*1024, backupCount=5, encoding='utf-8'),
        logging.StreamHandler()
    ]
)
LOGGER = logging.getLogger("nebula_ai")
LOGGER.info("Nebula AI starting up")
DEFAULT_MODEL: Final = "gemini-3.6-flash"
MODEL_ENV_VAR: Final = "ASKBUDDY_MODEL"
API_KEY_ENV_VARS: Final = ("GOOGLE_API_KEY", "GEMINI_API_KEY")
CHAT_FILE_TYPES: Final = [
    "pdf", "application/pdf", 
    "txt", "text/plain",
    "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "csv", "text/csv",
    "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "png", "image/png", "jpg", "image/jpeg", "jpeg",
    "md", "text/markdown"
]
STATEMENT_EXTENSIONS: Final = frozenset({"pdf", "txt", "csv", "xlsx", "png", "jpg", "jpeg", "webp"})
SPREADSHEET_EXTENSIONS: Final = frozenset({"csv", "xlsx"})
IMAGE_EXTENSIONS: Final = frozenset({"png", "jpg", "jpeg", "webp"})

SUGGESTIONS = {
    "Process a bank statement": "Process the attached bank statement and classify every transaction.",
    "Summarize spending": "Summarize spending by category from the processed statement.",
    "Export a recap": "Give me a short recap I can keep with the Excel export.",
    "Ask about a document": "What are the key findings in the uploaded file?",
}


def inject_shell_styles() -> None:
    """ChatGPT-like Nebula AI UI."""
    st.html(
        """
        <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
        
        html, body, [class*="css"] {
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif;
            font-size: 15px;
        }
        
        [data-testid="stHeader"], header, .stAppHeader {
            background: transparent !important;
            border: none !important;
        }
        .stApp {
            background: #0f0f0f !important;
        }
        .block-container {
            max-width: 48rem;
            padding-top: 0;
            padding-bottom: 9rem;
        }
        [data-testid="stSidebar"] {
            min-width: 17rem;
            max-width: 17rem;
            background: #171717 !important;
            border-right: 1px solid rgba(255, 255, 255, 0.08);
        }
        [data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p {
            color: #ececec;
        }
        [data-testid="stSidebar"] button {
            justify-content: flex-start;
            text-align: left;
            border-radius: 0.5rem;
            border: none;
            padding: 0.5rem 0.75rem;
            font-size: 0.875rem;
            color: #ececec;
            height: auto;
            line-height: 1.4;
        }
        [data-testid="stSidebar"] button:hover {
            background: #2a2a2a !important;
        }
        [data-testid="stSidebar"] hr {
            margin: 0.5rem 0;
            border-color: rgba(255, 255, 255, 0.06);
        }
        [data-testid="stSidebar"] [data-testid="stCaptionContainer"] {
            color: #8e8e8e;
            font-size: 0.75rem;
            text-transform: uppercase;
            letter-spacing: 0.05em;
        }
        [data-testid="stChatMessage"] {
            background: transparent;
            border: none;
            padding: 1.75rem 0;
        }
        [data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] {
            color: #ececec;
            line-height: 1.7;
        }
        [data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] p {
            margin-bottom: 1rem;
        }
        [data-testid="stChatInput"] {
            padding-bottom: 1.5rem;
        }
        [data-testid="stChatInput"] > div {
            border-radius: 1.5rem;
            border: 1px solid rgba(255, 255, 255, 0.12);
            background: #212121;
            box-shadow: none;
        }
        [data-testid="stChatInput"] textarea {
            font-size: 1rem;
            border-radius: 1.5rem;
            border: none;
            background: transparent;
            color: #ececec;
            padding: 0.8rem 1.25rem;
            line-height: 1.5;
        }
        [data-testid="stChatInput"] textarea::placeholder {
            color: #6b6b6b;
        }
        .hero-wrap {
            text-align: center;
            padding: 7rem 0 3rem;
            max-width: 42rem;
            margin: 0 auto;
        }
        .hero-brand {
            display: inline-flex;
            align-items: center;
            justify-content: center;
            margin-bottom: 1.5rem;
        }
        .hero-title {
            font-size: 2.25rem;
            font-weight: 600;
            letter-spacing: -0.02em;
            line-height: 1.15;
            margin: 0 0 0.75rem 0;
            color: #ececec;
        }
        .hero-sub {
            color: #9b9b9b;
            font-size: 1rem;
            line-height: 1.6;
        }
        .suggestion-btn {
            margin-bottom: 0.5rem;
        }
        .suggestion-btn button {
            background: #212121 !important;
            border: 1px solid rgba(255, 255, 255, 0.08) !important;
            border-radius: 0.75rem !important;
            padding: 0.875rem 1rem !important;
            font-size: 0.9rem !important;
            color: #ececec !important;
            text-align: left !important;
            justify-content: flex-start !important;
            height: auto !important;
            min-height: 3rem !important;
            line-height: 1.5 !important;
        }
        .suggestion-btn button:hover {
            background: #2a2a2a !important;
            border-color: rgba(255, 255, 255, 0.15) !important;
        }
        .nebula-brand {
            display: flex;
            align-items: center;
            gap: 0.625rem;
            padding: 0.5rem 0 1rem 0;
        }
        .nebula-brand-text {
            font-size: 1rem;
            font-weight: 600;
            color: #ececec;
            letter-spacing: -0.01em;
        }
        [data-testid="stVerticalBlock"] > [style*="flex-direction: column;"] > [data-testid="stVerticalBlock"] {
            gap: 0.5rem;
        }
        </style>
        """
    )


@st.cache_resource(show_spinner=False)
def get_processor() -> BankStatementProcessor:
    return BankStatementProcessor()


@st.cache_resource(show_spinner=False)
def get_ingestion() -> Any:
    return get_ingestion_service()


@st.cache_resource(show_spinner=False)
def warm_document_pipeline() -> bool:
    """Preload document models in the background.

    Docling spends ~100s loading its layout/table models on first use, and the
    embedding model a few more seconds. Paying that during the first document
    upload looks like a hang, so it happens once at server start instead.
    """
    def _warm() -> None:
        try:
            get_ingestion().warm()
            LOGGER.info("Document pipeline warm-up finished")
        except Exception as e:
            LOGGER.warning(f"Document pipeline warm-up failed: {e}")
        try:
            get_embedding_provider().get_embeddings().embed_query("warm up")
            LOGGER.info("Embedding model warm-up finished")
        except Exception as e:
            LOGGER.warning(f"Embedding warm-up failed: {e}")

    threading.Thread(target=_warm, name="nebula-warmup", daemon=True).start()
    return True


@st.cache_resource(show_spinner=False)
def get_vec_store() -> Any:
    return get_vector_store()


@st.cache_resource(show_spinner=False)
def get_doc_registry() -> Any:
    return get_document_registry()


@st.cache_resource(show_spinner=False)
def get_conv_store() -> Any:
    return get_conversation_store()


@st.cache_resource(show_spinner=False)
def get_orch() -> Any:
    return get_orchestrator()


def configured_api_key() -> str | None:
    """Return the configured Google API key without logging its value."""
    return next((os.getenv(name) for name in API_KEY_ENV_VARS if os.getenv(name)), None)


@st.cache_resource(show_spinner=False)
def create_llm(model_name: str, api_key: str) -> ChatGoogleGenerativeAI:
    """Create one reusable chat model instance for the Streamlit process."""
    return ChatGoogleGenerativeAI(
        model=model_name,
        google_api_key=api_key,
        temperature=0.2,
        max_retries=2,
        timeout=30,
    )


def initialise_session_state() -> None:
    """Initialize session state for chat, documents, and statements."""
    st.session_state.setdefault("messages", [])
    st.session_state.setdefault("document_record", None)
    st.session_state.setdefault("document_name", None)
    st.session_state.setdefault("document_loaded", False)
    st.session_state.setdefault("bank_statement", None)
    st.session_state.setdefault("pdf_password", "")
    st.session_state.setdefault("pending_prompt", "")
    # New multi-document RAG state
    st.session_state.setdefault("rag_documents", [])  # list of doc metadata
    st.session_state.setdefault("selected_documents", [])  # list of doc_ids
    st.session_state.setdefault("conversation_id", str(uuid.uuid4()))
    st.session_state.setdefault("use_web_search", False)
    st.session_state.setdefault("chat_history_list", [])  # list of conversations for sidebar
    st.session_state.setdefault("current_conversation_title", "New Chat")
    st.session_state.setdefault("document_summary", None)  # Summary shown after upload
    st.session_state.setdefault("auto_summary_done", {})  # Track if we showed summary


def clear_chat_state() -> None:
    """Clear chat, document, and statement context for new chat."""
    # Save current conversation if it has messages
    if st.session_state.get("messages"):
        save_current_conversation()
    
    st.session_state.messages = []
    st.session_state.document_record = None
    st.session_state.document_name = None
    st.session_state.document_loaded = False
    st.session_state.bank_statement = None
    st.session_state.pending_prompt = ""
    st.session_state.pop("_last_turn", None)
    st.session_state.conversation_id = str(uuid.uuid4())
    st.session_state.rag_documents = []
    st.session_state.selected_documents = []
    st.session_state.document_summary = None
    st.session_state.auto_summary_done = {}
    st.session_state.current_conversation_title = "New Chat"


def get_document_context(document_record: object) -> str:
    """Build a compact document-context string for prompt injection."""
    if document_record is None:
        return ""
    chunks = getattr(document_record, "chunks", [])
    text_parts = [chunk.text for chunk in chunks if getattr(chunk, "text", "").strip()]
    return "\n\n".join(text_parts)[:12000]


def get_statement_context(payload: dict[str, Any] | None) -> str:
    if not payload:
        return ""
    summary = payload.get("summary", {})
    rows = payload.get("transactions", [])[:80]
    lines = [
        f"Bank: {summary.get('bank_name')}",
        f"Holder: {summary.get('account_holder_name')}",
        f"Account: {summary.get('account_number')}",
        f"IFSC: {summary.get('ifsc_code')}",
        f"PDF type: {summary.get('document_type')}",
        f"Transactions: {summary.get('transaction_count')}",
    ]
    for row in rows:
        lines.append(
            f"{row.get('date')} | {row.get('description')} | "
            f"Dr {row.get('debit_amount')} Cr {row.get('credit_amount')} | "
            f"{row.get('category')}"
        )
    return "\n".join(lines)[:12000]


def stream_text(text: str) -> Iterator[str]:
    """Yield a text stream in small chunks to simulate flowing model output."""
    if not text:
        return
    pieces = re.findall(r".{1,180}\S*|.{1,180}", text)
    yield from pieces


def handle_uploaded_document(uploaded_file) -> None:
    """Persist extracted document metadata and content for later Q&A."""
    if uploaded_file is None:
        return
    settings = OllamaSettings.from_environment()
    service = DocumentAnalysisService(OllamaClient(settings), settings)
    try:
        record = service.ingest(uploaded_file.name, uploaded_file.getvalue())
    except DocumentAnalysisError as error:
        st.session_state.document_loaded = False
        st.session_state.document_record = None
        st.session_state.document_name = None
        st.session_state.document_error = str(error)
        return
    st.session_state.document_record = record
    st.session_state.document_name = uploaded_file.name
    st.session_state.document_loaded = True


def process_document_with_rag(uploaded_file, force_document: bool = False) -> dict[str, Any]:
    """Process document using new ingestion + RAG pipeline."""
    content = uploaded_file.getvalue()
    filename = uploaded_file.name
    
    validation = validate_upload(filename, content)
    if not validation.valid:
        return {"success": False, "error": validation.error}
    
    route = route_document(filename, content)
    
    # Bank statements go to specialized processor, unless the caller already
    # tried that route and wants this file ingested as a normal document.
    if not force_document and route.strategy == "bank_statement" and route.use_specialized:
        return {"success": False, "error": "BANK_STATEMENT", "route": route}
    
    try:
        ingestion = get_ingestion()
        registry = get_doc_registry()
        vec_store = get_vec_store()
        
        doc_id = str(uuid.uuid4())
        content_hash = hashlib.sha256(content).hexdigest()
        
        # Check if already indexed
        # (simple check - could be enhanced)
        
        # Ingest
        normalized, strategy = ingestion.ingest(filename, content)
        
        if normalized is None:
            return {"success": False, "error": "Could not process document"}
        
        # Register document
        reg_record = DocumentRecord(
            document_id=doc_id,
            filename=filename,
            file_type=validation.extension,
            file_size=len(content),
            content_hash=content_hash,
            session_id=st.session_state.get("conversation_id"),
            status="processing"
        )
        registry.register(reg_record)
        
        # Chunk
        chunker = StructureAwareChunker(chunk_size=1000, chunk_overlap=200)
        chunks = chunker.chunk_document(normalized)
        # Update chunk document ids
        for chunk in chunks:
            chunk.document_id = doc_id
        
        # Embed and store
        try:
            embedder = get_embedding_provider()
            embeddings = embedder.get_embeddings()
            chunk_texts = [c.text for c in chunks]
            chunk_embeddings = embeddings.embed_documents(chunk_texts)
            vec_store.add_chunks(chunks, chunk_embeddings)
            registry.update_status(doc_id, "indexed", chunk_count=len(chunks))
        except Exception as e:
            registry.update_status(doc_id, "error", error_message=str(e))
            return {"success": False, "error": f"Embedding/indexing failed: {e}"}
        
        doc_info = {
            "document_id": doc_id,
            "filename": filename,
            "file_type": validation.extension,
            "chunk_count": len(chunks),
            "status": "indexed"
        }
        # Check if already exists
        existing_ids = [d['document_id'] for d in st.session_state.rag_documents]
        if doc_id not in existing_ids:
            st.session_state.rag_documents.append(doc_info)
        if doc_id not in st.session_state.selected_documents:
            st.session_state.selected_documents.append(doc_id)
        else:
            # Ensure it's selected
            pass
        
        return {"success": True, "document": doc_info}
        
    except Exception as e:
        return {"success": False, "error": str(e)}


def statement_payload(statement: BankStatement) -> dict[str, Any]:
    return {
        "summary": statement.as_summary(),
        "warnings": statement.warnings,
        "transactions": [item.as_record() for item in statement.transactions],
    }


def payload_to_transactions(payload: dict[str, Any]) -> list[Transaction]:
    return [Transaction(**row) for row in payload.get("transactions", [])]


def file_extension(name: str) -> str:
    """Return the lower-cased extension of a filename (``""`` when absent)."""
    return name.lower().rsplit(".", 1)[-1] if "." in name else ""


def read_statement_rows(name: str, data: bytes) -> list[list[object]]:
    """Read CSV/XLSX cells into rows for tabular statement processing."""
    extension = file_extension(name)
    if extension == "csv":
        import csv
        from io import StringIO

        return [list(row) for row in csv.reader(StringIO(data.decode("utf-8-sig", errors="replace")))]
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    rows: list[list[object]] = []
    for sheet in workbook.worksheets:
        rows.extend(list(row) for row in sheet.iter_rows(values_only=True))
    return rows


def process_bank_file(uploaded_file, password: str = "") -> BankStatement:
    """Route any supported statement upload (PDF/CSV/XLSX/TXT/image) to the processor."""
    processor = get_processor()
    name = uploaded_file.name
    data = uploaded_file.getvalue()
    extension = file_extension(name)
    if extension == "pdf" or b"%PDF" in data[:1024]:
        return processor.process_pdf(data, name, password=password)
    if extension in SPREADSHEET_EXTENSIONS:
        return processor.process_tabular(read_statement_rows(name, data), name)
    if extension in IMAGE_EXTENSIONS:
        return processor.process_image(data, name)
    return processor.process_text(data.decode("utf-8", errors="ignore"), name)


def read_bank_statement(uploaded_file, password: str = "") -> BankStatement:
    """Process a statement upload with visible progress.

    Scanned statements are OCR'd locally, which takes tens of seconds per page;
    without a status widget the app looks frozen because nothing is rendered
    until the whole parse finishes.
    """
    name = uploaded_file.name
    with st.status(f"Reading {name}…", expanded=False) as status:
        status.update(label=f"Reading {name}… extracting text (OCR may take a minute)")
        try:
            statement = process_bank_file(uploaded_file, password=password)
        except Exception:
            status.update(label=f"Could not read {name}", state="error")
            raise
        status.update(
            label=f"Read {name} — {len(statement.transactions)} transaction(s)",
            state="complete",
            expanded=False,
        )
        return statement


def save_current_conversation() -> None:
    """Save current conversation to persistent store."""
    try:
        conv_store = get_conv_store()
        messages = st.session_state.get("messages", [])
        if not messages:
            return
        conv_id = st.session_state.get("conversation_id")
        if not conv_id:
            return
        # Create conversation if not exists
        conv = Conversation(
            conversation_id=conv_id,
            title=st.session_state.get("current_conversation_title", "New Chat")[:60],
            created_at=datetime.now().timestamp(),
            updated_at=datetime.now().timestamp()
        )
        conv_store.create_conversation(conv)
        # Save messages
        for msg in messages:
            msg_id = str(uuid.uuid4())
            conv_msg = Message(
                message_id=msg_id,
                conversation_id=conv_id,
                role=msg["role"],
                content=msg["content"],
                metadata={"has_statement": bool(msg.get("statement")), "has_sources": bool(msg.get("sources"))}
            )
            conv_store.add_message(conv_msg)
    except Exception as e:
        pass  # Don't fail on save errors


def load_conversation(conv_id: str) -> None:
    """Load a conversation from persistent store."""
    try:
        conv_store = get_conv_store()
        messages = conv_store.get_conversation_messages(conv_id)
        st.session_state.messages = []
        for m in messages:
            st.session_state.messages.append({
                "role": m.role,
                "content": m.content,
            })
        st.session_state.conversation_id = conv_id
        # Try to get title
        convs = conv_store.list_conversations()
        for c in convs:
            if c.conversation_id == conv_id:
                st.session_state.current_conversation_title = c.title
                break
        st.session_state.document_summary = None
        st.session_state.bank_statement = None
        st.session_state.rag_documents = []
        st.session_state.selected_documents = []
        st.rerun()
    except Exception as e:
        st.error(f"Failed to load conversation: {e}")


def render_sidebar() -> None:
    """Refined Nebula AI sidebar."""
    with st.sidebar:
        st.html(
            """
            <div class="nebula-brand">
                <svg width="20" height="20" viewBox="0 0 24 24" fill="none" xmlns="http://www.w3.org/2000/svg">
                    <path d="M12 2L14.5 9.5H22L16 14L18.5 21.5L12 17L5.5 21.5L8 14L2 9.5H9.5L12 2Z" fill="url(#grad)" />
                    <defs>
                        <linearGradient id="grad" x1="2" y1="2" x2="22" y2="21.5" gradientUnits="userSpaceOnUse">
                            <stop stop-color="#8b5cf6"/>
                            <stop offset="1" stop-color="#06b6d4"/>
                        </linearGradient>
                    </defs>
                </svg>
                <span class="nebula-brand-text">Nebula AI</span>
            </div>
            """
        )
        
        if st.button("+ New chat", use_container_width=True, key="new_chat_main"):
            save_current_conversation()
            clear_chat_state()
            st.rerun()

        st.markdown("---")
        st.caption("CONVERSATIONS")
        try:
            conv_store = get_conv_store()
            conversations = conv_store.list_conversations()
            if conversations:
                for conv in conversations[:15]:
                    title = (conv.title or "New Chat").strip()
                    if len(title) > 28:
                        title = title[:25] + "..."
                    is_current = conv.conversation_id == st.session_state.get("conversation_id")
                    if st.button(title, use_container_width=True, key=f"conv_{conv.conversation_id}", type="secondary" if not is_current else "primary"):
                        load_conversation(conv.conversation_id)
            else:
                st.caption("No conversations yet")
        except Exception:
            pass

        st.markdown("---")
        st.caption("DOCUMENTS")
        if st.session_state.rag_documents or st.session_state.get("bank_statement"):
            if st.session_state.rag_documents:
                for doc in st.session_state.rag_documents:
                    col1, col2 = st.columns([4, 1])
                    with col1:
                        is_sel = doc['document_id'] in st.session_state.selected_documents
                        fname = doc['filename']
                        if len(fname) > 22:
                            fname = fname[:19] + "..."
                        if st.checkbox(fname, value=is_sel, key=f"sel_{doc['document_id']}"):
                            if doc['document_id'] not in st.session_state.selected_documents:
                                st.session_state.selected_documents.append(doc['document_id'])
                        else:
                            if doc['document_id'] in st.session_state.selected_documents:
                                st.session_state.selected_documents = [d for d in st.session_state.selected_documents if d != doc['document_id']]
                    with col2:
                        if st.button("×", key=f"del_{doc['document_id']}", help="Remove"):
                            try:
                                get_vec_store().delete_document(doc['document_id'])
                                get_doc_registry().delete(doc['document_id'])
                                st.session_state.rag_documents = [d for d in st.session_state.rag_documents if d['document_id'] != doc['document_id']]
                                st.session_state.selected_documents = [d for d in st.session_state.selected_documents if d != doc['document_id']]
                                st.rerun()
                            except Exception as e:
                                LOGGER.error(f"Delete failed: {e}")
            if st.session_state.get("bank_statement"):
                fname = st.session_state.bank_statement["summary"].get("file_name") or "Bank Statement"
                if len(fname) > 22:
                    fname = fname[:19] + "..."
                st.caption(fname)
                st.success("Loaded", icon=":material/check_circle:")
        else:
            st.caption("No documents")

        st.markdown("---")
        with st.expander("Settings"):
            st.text_input(
                "PDF password",
                type="password",
                key="pdf_password",
                help="For password-protected PDFs",
            )
            st.checkbox("Enable web search", key="use_web_search")


def render_top_bar() -> None:
    """Render top bar with document indicators."""
    has_docs = bool(st.session_state.rag_documents) or bool(st.session_state.get("bank_statement"))
    if has_docs:
        cols = st.columns([4, 1])
        with cols[0]:
            doc_parts = []
            for doc in st.session_state.rag_documents:
                doc_parts.append(f"📄 {doc['filename'][:25]}")
            if st.session_state.get("bank_statement"):
                doc_parts.append(f"🏦 {st.session_state.bank_statement['summary'].get('file_name','Statement')[:25]}")
            st.markdown(" • ".join(doc_parts))
        with cols[1]:
            if st.button("Clear", key="clear_docs_top", help="Clear all documents"):
                try:
                    for doc in st.session_state.rag_documents:
                        get_vec_store().delete_document(doc['document_id'])
                        get_doc_registry().delete(doc['document_id'])
                except Exception:
                    pass
                st.session_state.rag_documents = []
                st.session_state.selected_documents = []
                st.session_state.bank_statement = None
                st.session_state.document_record = None
                st.session_state.document_loaded = False
                st.rerun()
    else:
        st.markdown("")


def render_empty_state() -> None:
    """Nebula AI welcome state."""
    render_top_bar()
    st.html(
        """
        <div class="hero-wrap">
            <div class="hero-brand">
                <svg width="32" height="32" viewBox="0 0 24 24" fill="none" xmlns="http://www.w3.org/2000/svg">
                    <path d="M12 2L14.5 9.5H22L16 14L18.5 21.5L12 17L5.5 21.5L8 14L2 9.5H9.5L12 2Z" fill="url(#grad4)" />
                    <defs>
                        <linearGradient id="grad4" x1="2" y1="2" x2="22" y2="21.5" gradientUnits="userSpaceOnUse">
                            <stop stop-color="#8b5cf6"/>
                            <stop offset="1" stop-color="#06b6d4"/>
                        </linearGradient>
                    </defs>
                </svg>
            </div>
            <p class="hero-title">How can I help you?</p>
            <p class="hero-sub">Upload documents, analyze bank statements, or ask me anything.</p>
        </div>
        """
    )
    labels = list(SUGGESTIONS.keys())
    cols = st.columns(2, gap="medium")
    for index, label in enumerate(labels):
        with cols[index % 2]:
            st.markdown('<div class="suggestion-btn">', unsafe_allow_html=True)
            if st.button(label, use_container_width=True, key=f"suggest_{index}"):
                st.session_state.pending_prompt = SUGGESTIONS[label]
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)


def render_statement_card(payload: dict[str, Any], key_prefix: str) -> None:
    summary = payload.get("summary", {})
    rows = payload.get("transactions", [])
    cols = st.columns(4)
    cols[0].metric("Transactions", summary.get("transaction_count") or 0)
    cols[1].metric("Debits", f"{float(summary.get('total_debit') or 0):,.2f}")
    cols[2].metric("Credits", f"{float(summary.get('total_credit') or 0):,.2f}")
    cols[3].metric("PDF type", str(summary.get("document_type") or "text").title())

    st.markdown(
        f"**{summary.get('account_holder_name') or 'Account holder unknown'}**  \n"
        f"Account `{summary.get('account_number') or '—'}` · IFSC `{summary.get('ifsc_code') or '—'}`  \n"
        f"{summary.get('bank_name') or 'Bank not detected'}"
    )
    for warning in payload.get("warnings") or []:
        st.caption(warning)

    if rows:
        st.dataframe(rows, width="stretch", hide_index=True)
    else:
        st.info("The PDF was read, but no transaction rows matched this layout.")

    processor = get_processor()
    transactions = payload_to_transactions(payload)
    statement = BankStatement(
        file_name=str(summary.get("file_name") or "statement.pdf"),
        document_type=str(summary.get("document_type") or "text"),
        bank_name=summary.get("bank_name"),
        account_holder_name=summary.get("account_holder_name"),
        account_number=summary.get("account_number"),
        ifsc_code=summary.get("ifsc_code"),
        transactions=transactions,
        warnings=list(payload.get("warnings") or []),
    )
    export_cols = st.columns(2)
    with export_cols[0]:
        st.download_button(
            "Download CSV",
            data=processor.export_bytes(transactions, "csv", statement),
            file_name="bank_statement.csv",
            mime="text/csv",
            icon=":material/download:",
            width="stretch",
            key=f"{key_prefix}_csv",
        )
    with export_cols[1]:
        st.download_button(
            "Download Excel",
            data=processor.export_bytes(transactions, "xlsx", statement),
            file_name="bank_statement.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            icon=":material/download:",
            width="stretch",
            key=f"{key_prefix}_xlsx",
        )


def render_messages() -> None:
    """Render chat history."""
    render_top_bar()
    for index, message in enumerate(st.session_state.messages):
        avatar = ":material/person:" if message["role"] == "user" else ":material/auto_awesome:"
        with st.chat_message(message["role"], avatar=avatar):
            if message.get("content"):
                st.markdown(message["content"])
            payload = message.get("statement")
            if payload:
                render_statement_card(payload, key_prefix=f"msg_{index}")
            sources = message.get("sources")
            if sources:
                with st.expander("Sources"):
                    for s in sources:
                        if s.get("type") == "web":
                            st.markdown(f"- [{s.get('title')}]({s.get('url')})")
                        elif s.get("type") == "document":
                            src = f"- {s.get('filename', 'Unknown')}"
                            if s.get('page'):
                                src += f" (Page {s.get('page')})"
                            st.markdown(src)
                        else:
                            st.markdown(f"- {s.get('filename', s.get('title', 'Source'))}")


def build_prompt(question: str) -> str:
    """Create a grounded prompt that uses statement or document context when present."""
    base = (
        "You are AskBuddy, a helpful AI assistant. Answer clearly and concisely. "
        "When bank-statement context is present, use those extracted rows. "
        "Do not reclassify transactions; categories are already assigned by a non-LLM engine. "
        "If documents are provided, ground your answers in the document content. "
        "Maintain conversation context."
    )
    statement_context = get_statement_context(st.session_state.get("bank_statement"))
    if statement_context:
        return f"{base}\n\nBANK STATEMENT CONTEXT:\n{statement_context}\n\nUSER QUESTION:\n{question}"
    if st.session_state.document_record is not None:
        context = get_document_context(st.session_state.document_record)
        if context:
            return f"{base}\n\nDOCUMENT CONTEXT:\n{context}\n\nUSER QUESTION:\n{question}"
    return f"{base}\n\nUSER QUESTION:\n{question}"


def generate_response(prompt: str) -> str:
    """Generate the final answer text from the configured Gemini model."""
    model_name = os.getenv(MODEL_ENV_VAR, DEFAULT_MODEL)
    llm = create_llm(model_name, configured_api_key())
    return ask_model(llm, prompt)


def consume_prompt(prompt: Any) -> tuple[str, list[Any]]:
    if prompt is None:
        return "", []
    if isinstance(prompt, str):
        return prompt.strip(), []
    text = (getattr(prompt, "text", None) or "").strip()
    files = list(getattr(prompt, "files", None) or [])
    return text, files


def is_statement_file(uploaded_file) -> bool:
    """True when the upload looks like a bank statement rather than a general document."""
    name = uploaded_file.name
    data = uploaded_file.getvalue()
    return file_extension(name) in STATEMENT_EXTENSIONS or b"%PDF" in data[:1024]


def generate_document_summary(filename: str, document_ids: list[str]) -> str:
    """Generate an initial summary of uploaded document."""
    try:
        orch = get_orch()
        question = f"Provide a clear and comprehensive summary of the document '{filename}'. Include: overview, key points, main topics, and important details. Keep it concise but informative."
        result = orch.orchestrate(
            question=question,
            document_ids=document_ids,
            has_documents=True,
            conversation_id=st.session_state.conversation_id
        )
        if result and result.answer and result.answer.strip():
            return result.answer
    except Exception:
        pass
    
    if configured_api_key():
        try:
            llm = create_llm(os.getenv(MODEL_ENV_VAR, DEFAULT_MODEL), configured_api_key())
            prompt = f"Summarize this document '{filename}' comprehensively. Provide an overview, key points, and main topics. Be clear and concise."
            answer = ask_model(llm, prompt)
            if answer and answer.strip():
                return answer
        except Exception:
            pass
    
    return f"**Document uploaded:** {filename}\n\nI've indexed this document and you can now ask me any questions about it."


def handle_turn(user_text: str, files: list[Any]) -> None:
    display = user_text or ("Uploaded " + ", ".join(item.name for item in files))
    if len(st.session_state.messages) == 0 and user_text:
        title = user_text[:50].strip()
        if title:
            st.session_state.current_conversation_title = title + ("..." if len(user_text) > 50 else "")
    elif len(st.session_state.messages) == 0 and files and not user_text:
        st.session_state.current_conversation_title = f"Chat about {files[0].name}"[:50]
    
    st.session_state.messages.append({"role": "user", "content": display})
    with st.chat_message("user", avatar=":material/person:"):
        st.markdown(display)
        for item in files:
            st.caption(f"📎 {item.name}")

    statement_payload_data = None
    status_notes: list[str] = []
    newly_uploaded_docs = []
    
    # Process files
    for item in files:
        data = item.getvalue()
        statement_attempted = False

        # Only uploads that actually look like statements go through the
        # specialized parser. Parsing every PDF as a statement used to make
        # ordinary documents come back as a 0-transaction card, so they never
        # reached the document/RAG pipeline and nothing was generated.
        try:
            route = route_document(item.name, data)
            looks_like_statement = route.use_specialized and route.strategy == "bank_statement"
        except Exception as e:
            LOGGER.warning(f"Could not route {item.name}: {e}")
            looks_like_statement = False

        if looks_like_statement:
            statement_attempted = True
            statement = None
            try:
                statement = read_bank_statement(item, password=str(st.session_state.get("pdf_password") or ""))
            except Exception as e:
                LOGGER.info(f"{item.name}: not a usable bank statement ({e}); using the document pipeline")
            if statement is not None and statement.transactions:
                statement_payload_data = statement_payload(statement)
                st.session_state.bank_statement = statement_payload_data
                status_notes.append(
                    f"Processed {item.name} as bank statement ({len(statement.transactions)} transactions)."
                )
                continue
            if statement is not None:
                LOGGER.info(f"{item.name}: statement parse found no transactions; using the document pipeline")

        # Process as general document
        try:
            with st.spinner(f"Reading and indexing {item.name}…"):
                result = process_document_with_rag(item, force_document=statement_attempted)
                if result.get("success"):
                    newly_uploaded_docs.append(result["document"])
                    status_notes.append(f"Uploaded {item.name}")
                    continue
                if result.get("error") == "BANK_STATEMENT":
                    try:
                        statement = read_bank_statement(item, password=str(st.session_state.get("pdf_password") or ""))
                    except Exception as e:
                        status_notes.append(f"Error: {e}")
                    else:
                        if statement.transactions:
                            statement_payload_data = statement_payload(statement)
                            st.session_state.bank_statement = statement_payload_data
                            status_notes.append(
                                f"Processed {item.name} as bank statement ({len(statement.transactions)} transactions)."
                            )
                            continue
                        LOGGER.info(f"{item.name}: statement parse found no transactions; using the document pipeline")
                    result = process_document_with_rag(item, force_document=True)
                    if result.get("success"):
                        newly_uploaded_docs.append(result["document"])
                        status_notes.append(f"Uploaded {item.name}")
                        continue
                handle_uploaded_document(item)
                if st.session_state.get("document_error"):
                    status_notes.append(st.session_state.document_error)
                    st.session_state.document_error = ""
                elif st.session_state.get("document_loaded"):
                    status_notes.append(f"Processed {item.name}")
        except Exception as e:
            LOGGER.error(f"Error processing {item.name}: {e}")
            status_notes.append(f"Error: {e}")

    question = user_text.strip() if user_text else ""
    assistant_sources = []
    assistant_text = ""
    
    with st.chat_message("assistant", avatar=":material/auto_awesome:"):
        if statement_payload_data:
            if status_notes:
                st.markdown("\n".join(status_notes))
            render_statement_card(statement_payload_data, key_prefix="live")
            assistant_text = "\n".join(status_notes) if status_notes else "Statement processed."
        elif newly_uploaded_docs:
            if status_notes:
                st.markdown("\n".join(status_notes))
            st.markdown("---")
            st.markdown("**Document Summary**")
            for doc in newly_uploaded_docs:
                doc_ids = [doc['document_id']]
                if doc['document_id'] not in st.session_state.selected_documents:
                    st.session_state.selected_documents.append(doc['document_id'])
                with st.spinner(f"Summarizing {doc['filename']}..."):
                    try:
                        summary = generate_document_summary(doc['filename'], doc_ids)
                        if summary and summary.strip():
                            st.markdown(summary)
                            assistant_text += summary + "\n\n"
                        else:
                            msg = f"I've uploaded **{doc['filename']}**. You can ask me anything about this document."
                            st.markdown(msg)
                            assistant_text += msg + "\n\n"
                    except Exception as e:
                        msg = f"I've uploaded **{doc['filename']}**. You can ask me anything about this document."
                        st.markdown(msg)
                        assistant_text += msg + "\n\n"
            # If user asked a question with the upload
            if question and question.strip() and question.lower() not in ["summarize this document.", "process this file."]:
                st.markdown("---")
                st.markdown("**Answer**")
                doc_ids = st.session_state.selected_documents if st.session_state.selected_documents else [d['document_id'] for d in newly_uploaded_docs]
                with st.spinner("Answering..."):
                    try:
                        orch = get_orch()
                        result = orch.orchestrate(
                            question=question,
                            document_ids=doc_ids,
                            has_documents=True,
                            conversation_id=st.session_state.conversation_id
                        )
                        if result and result.answer and result.answer.strip():
                            st.markdown(result.answer)
                            assistant_text += "\n\n" + result.answer
                            assistant_sources = getattr(result, 'sources', []) or []
                        elif configured_api_key():
                            try:
                                ans = generate_response(build_prompt(question))
                                st.write_stream(stream_text(ans))
                                assistant_text += "\n\n" + ans
                            except Exception as e2:
                                st.error(str(e2))
                    except Exception as e:
                        if configured_api_key():
                            try:
                                ans = generate_response(build_prompt(question))
                                st.write_stream(stream_text(ans))
                                assistant_text += "\n\n" + ans
                            except Exception as e2:
                                st.error(str(e2))
        elif st.session_state.rag_documents:
            # Answer based on documents - use document context first, then web if needed
            doc_ids = st.session_state.selected_documents if st.session_state.selected_documents else [d['document_id'] for d in st.session_state.rag_documents]
            if not st.session_state.selected_documents:
                st.session_state.selected_documents = doc_ids
            q = question if question else "Answer based on the uploaded documents."
            with st.spinner("Answering..."):
                try:
                    orch = get_orch()
                    result = orch.orchestrate(
                        question=q,
                        document_ids=doc_ids,
                        has_documents=True,
                        conversation_id=st.session_state.conversation_id
                    )
                    if result and result.answer and result.answer.strip():
                        st.markdown(result.answer)
                        assistant_text = result.answer
                        assistant_sources = getattr(result, 'sources', []) or []
                    else:
                        if configured_api_key():
                            ans = generate_response(build_prompt(q))
                            st.write_stream(stream_text(ans))
                            assistant_text = ans
                        else:
                            assistant_text = "I need more context to answer. Please upload documents or try again."
                            st.markdown(assistant_text)
                except Exception as e:
                    if configured_api_key():
                        try:
                            ans = generate_response(build_prompt(q))
                            st.write_stream(stream_text(ans))
                            assistant_text = ans
                        except Exception as e2:
                            st.error(str(e2))
                            assistant_text = str(e2)
                    else:
                        st.error("Error processing request")
                        assistant_text = "Error processing request"
        elif question:
            if configured_api_key():
                try:
                    ans = generate_response(build_prompt(question))
                    st.write_stream(stream_text(ans))
                    assistant_text = ans
                except Exception as error:
                    st.error(str(error))
                    assistant_text = str(error)
            else:
                assistant_text = "Set GOOGLE_API_KEY to chat or upload documents."
                st.markdown(assistant_text)
        else:
            assistant_text = "How can I help you?"
            st.markdown(assistant_text)

    st.session_state.messages.append(
        {
            "role": "assistant",
            "content": assistant_text if not statement_payload_data else ("\n".join(status_notes) if status_notes else assistant_text),
            "statement": statement_payload_data,
            "sources": assistant_sources if assistant_sources else None,
        }
    )
    try:
        save_current_conversation()
    except Exception:
        pass


def main() -> None:
    """Render the Nebula AI workspace."""
    st.set_page_config(
        page_title="Nebula AI - Turn Documents into Knowledge",
        page_icon="✨",
        layout="centered",
        initial_sidebar_state="expanded",
    )
    inject_shell_styles()
    initialise_session_state()
    warm_document_pipeline()
    render_sidebar()

    if not st.session_state.messages:
        render_empty_state()
    else:
        render_messages()

    prompt = st.chat_input(
        "Message Nebula AI...",
        accept_file="multiple",
        file_type=CHAT_FILE_TYPES,
        max_upload_size=50,
    )

    pending = st.session_state.get("pending_prompt") or ""
    if pending and prompt is None:
        st.session_state.pending_prompt = ""
        handle_turn(pending, [])
        st.rerun()

    user_text, files = consume_prompt(prompt)
    if user_text or files:
        signature = (user_text, tuple(sorted((item.name, item.size) for item in files)))
        if st.session_state.get("_last_turn") != signature:
            handle_turn(user_text, files)
            # Guards the script rerun that follows a submit, so the same upload
            # is not processed twice.
            st.session_state._last_turn = signature
            st.rerun()
    else:
        # No submission in this run: forget the signature so that uploading the
        # same file again is not silently ignored.
        st.session_state.pop("_last_turn", None)


if __name__ == "__main__":
    main()
